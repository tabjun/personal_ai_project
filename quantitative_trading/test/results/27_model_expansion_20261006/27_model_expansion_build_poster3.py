# -*- coding: utf-8 -*-
"""계명대 학술 포스터 3 — 비정상 시계열 변동성 예측 (90x120cm)"""
import os, re
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import ImageFont, Image

NAVY_BG = RGBColor(0x1F, 0x4E, 0x79)
BAR     = RGBColor(0x20, 0x38, 0x64)
LTBLUE  = RGBColor(0xDE, 0xEA, 0xF6)
LTBLUE2 = RGBColor(0xF2, 0xF7, 0xFC)
EDGE    = RGBColor(0x9D, 0xC3, 0xE6)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
INK     = RGBColor(0x1A, 0x1A, 0x1A)
GREY    = RGBColor(0x55, 0x5F, 0x6A)
RED     = RGBColor(0xC0, 0x00, 0x00)
PALE    = RGBColor(0xCF, 0xE0, 0xEE)
GOLD    = RGBColor(0xFF, 0xD9, 0x66)
SEQ     = RGBColor(0x1F, 0x6F, 0xB2)   # 순차 = 파랑
PARF    = RGBColor(0xE0, 0x7B, 0x39)   # 병렬 특징 기반 = 주황
PARD    = RGBColor(0x8C, 0x5A, 0x2B)   # 병렬 딥러닝 = 갈색
FONT    = "맑은 고딕"

FS   = float(os.environ.get("FS", "1.0"))
WF1  = float(os.environ.get("WF1", "85.4"))
WF2  = float(os.environ.get("WF2", "64.0"))
WF3  = float(os.environ.get("WF3", "41.9"))
def s_(x): return round(x * FS, 1)

FPATH = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
_c = {}
def _f(px):
    k = int(px)
    if k not in _c: _c[k] = ImageFont.truetype(FPATH, k)
    return _c[k]
def _w(t, pt): return _f(round(pt * 4)).getlength(t) / 4 / 72 * 2.54

SAFE = 0.92
def wrap(text, pt, width_cm):
    width_cm *= SAFE
    out = []
    for para in text.split("\n"):
        if not para:
            out.append(""); continue
        toks, buf = [], ""
        for ch in para:
            if ch.isascii() and (ch.isalnum() or ch in "+-.%/±−~,"):
                buf += ch
            else:
                if buf: toks.append(buf); buf = ""
                toks.append(ch)
        if buf: toks.append(buf)
        line = ""
        for t in toks:
            cand = line + t
            if line and _w(cand.rstrip(), pt) > width_cm:
                out.append(line.rstrip()); line = "" if t == " " else t
            else: line = cand
        out.append(line.rstrip())
    return out

LH = 1.33
def th(text, pt, w, spacing=1.0, extra=0.0):
    return len(wrap(_plain(text), pt, w)) * pt * LH * spacing / 72 * 2.54 + extra + 0.08

def _plain(t):
    return t.replace("~", "").replace("^", "") if isinstance(t, str) else t

prs = Presentation("base.pptx")
sld = prs.slides[0]
for sh in list(sld.shapes): sh._element.getparent().remove(sh._element)

def _noautofit(tf):
    bp = tf._txBody.bodyPr
    for tag in ("a:normAutofit", "a:spAutoFit"):
        for e in bp.findall(qn(tag)): bp.remove(e)
    bp.append(bp.makeelement(qn("a:noAutofit"), {}))

def box(x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, adj=None):
    s = sld.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    if adj is not None:
        try: s.adjustments[0] = adj
        except Exception: pass
    s.text_frame.text = ""
    return s

def _style(r, size, bold, color, base=0):
    f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    if base: rPr.set("baseline", str(base))
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None: e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", FONT)

TOK = re.compile(r"(~[^~]+~|\^[^\^]+\^)")
def _emit(para, t, size, bold, color):
    """~x~ = 아래첨자, ^x^ = 위첨자"""
    for piece in TOK.split(t):
        if not piece: continue
        if piece.startswith("~") and piece.endswith("~") and len(piece) > 2:
            r = para.add_run(); r.text = piece[1:-1]; _style(r, size, bold, color, -25000)
        elif piece.startswith("^") and piece.endswith("^") and len(piece) > 2:
            r = para.add_run(); r.text = piece[1:-1]; _style(r, size, bold, color, 30000)
        else:
            r = para.add_run(); r.text = piece; _style(r, size, bold, color)

def text(x, y, w, h, runs, size=19, color=INK, bold=False, align=PP_ALIGN.LEFT,
         spacing=1.0, anchor=MSO_ANCHOR.TOP):
    tb = sld.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame; tf.word_wrap = True; _noautofit(tf)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.line_spacing = spacing
        segs = p if isinstance(p, list) else [(p, {})]
        for t, o in segs:
            _emit(para, t, o.get("size", size), o.get("bold", bold), o.get("color", color))
    return tb

def pic(path, x, y, w):
    return sld.shapes.add_picture(path, Cm(x), Cm(y), Cm(w))
def ar(p):
    w, h = Image.open(p).size
    return h / w

def secbar(x, y, w, label, h=2.3, size=None):
    box(x, y, w, h, BAR, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, adj=0.16)
    text(x, y, w, h, label, size=size or s_(36), color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def rbox(x, y, w, title, items, isize, tsize=None, pad=0.42, gap=0.12,
         fill=LTBLUE, tcolor=RED, icolor=INK, ibold=False, spacing=1.06):
    """둥근 상자 + 제목 + 항목들. 높이를 돌려준다."""
    tsize = tsize or isize + 1
    iw = w - 1.8
    hts = [th(it if isinstance(it, str) else "".join(a for a, _ in it), isize, iw,
              spacing) for it in items]
    ht = th(title, tsize, iw) if title else 0.0
    H = pad + (ht + gap if title else 0) + sum(hts) + gap * (len(items) - 1) + pad
    box(x, y, w, H, fill, EDGE, 1.3, MSO_SHAPE.ROUNDED_RECTANGLE, 0.10)
    yy = y + pad
    if title:
        text(x + 0.9, yy, iw, ht, title, size=tsize, color=tcolor, bold=True)
        yy += ht + gap
    for it, hh in zip(items, hts):
        text(x + 0.9, yy, iw, hh, [it] if isinstance(it, list) else it,
             size=isize, color=icolor, bold=ibold, spacing=spacing)
        yy += hh + gap
    return H

def table(x, y, w, rows, colw, hsize, bsize, rowpad=0.34, aligns=None, bolds=None):
    tot = sum(colw); cw = [w * c / tot for c in colw]
    hs = []
    for ri, row in enumerate(rows):
        sz = hsize if ri == 0 else bsize
        hs.append(max(th(c[0] if isinstance(c, tuple) else str(c), sz, cw[ci] - 0.45)
                      for ci, c in enumerate(row)) + rowpad)
    gf = sld.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(sum(hs)))
    tb = gf.table; tb.first_row = True; tb.horz_banding = False
    for ci, c in enumerate(tb.columns): c.width = Cm(cw[ci])
    for ri, r in enumerate(tb.rows): r.height = Cm(hs[ri])
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tb.cell(ri, ci)
            cell.margin_left = cell.margin_right = Cm(0.20)
            cell.margin_top = cell.margin_bottom = Cm(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = BAR if ri == 0 else (WHITE if ri % 2 else LTBLUE2)
            tf = cell.text_frame; tf.word_wrap = True; _noautofit(tf)
            p = tf.paragraphs[0]; p.line_spacing = 1.0
            if aligns: p.alignment = aligns[ci]
            s, col = (val if isinstance(val, tuple) else (str(val), None))
            if col is None: col = WHITE if ri == 0 else INK
            bold = ri == 0 or (bolds and ci in bolds) or isinstance(val, tuple)
            _emit(p, s, hsize if ri == 0 else bsize, bold, col)
    return sum(hs)

# ================================================================ 치수
ML, MR = 2.3, 87.7
FW = MR - ML
CW = (FW - 1.6) / 2
LX, RX = ML, ML + CW + 1.6
PANEL_T, PANEL_B = 5.2, 119.2
B   = s_(19)     # 본문 18~20
CAP = 15.0       # 캡션 15~16 (하한 고정)
TB  = 15.0
TH_ = 15.0
BX  = B
GAP, CGAP = float(os.environ.get("GAP", "0.50")), 0.20

# ---------------- 머리 (60mm 이내)
sld.shapes.add_picture("kmu_logo.png", Cm(1.9), Cm(0.45), Cm(4.3))
text(8.0, 0.12, 74.0, 2.25, "비정상 시계열 변동성 예측",
     size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(8.0, 2.28, 74.0, 1.45,
     [[("순차", dict(color=RGBColor(0x8E, 0xC6, 0xF5))), (" vs ", {}),
       ("병렬", dict(color=GOLD)), (" 알고리즘 성능 비교", {})]],
     size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(8.0, 3.80, 74.0, 1.15,
     "윤태준* · 김채원* · 손낙훈*   |   *계명대학교 자연과학대학 통계학과",
     size=21, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

box(0.8, PANEL_T, 88.4, PANEL_B - PANEL_T, WHITE,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.012)

y0 = 5.9
secbar(LX, y0, CW, "연구배경 및 연구목적")
secbar(RX, y0, CW, "분석자료 및 분석방법")

# ================================================================ 왼쪽 단
y = y0 + 2.3 + 0.55
intro = "암호화폐는 24시간 거래되고, 크게 움직인 뒤에는 계속 크게 움직인다(변동성 군집)."
h = th(intro, B, CW); text(LX, y, CW, h, intro, size=B, color=BAR, bold=True,
                           spacing=1.06); y += h + 0.18
BUL = ["가격은 제자리로 돌아오지 않는다: 20종목 중 18종목이 단위근(ADF · KPSS 검정).",
       "수익률은 평균은 안정적이지만 분산이 시간에 따라 변한다(ARCH-LM 검정, 20/20종목).",
       "변동성은 단위근은 없지만 높거나 낮은 상태가 오래 이어진다(정상성 판정이 엇갈린 종목: "
       "15분 19/20 → 12시간 11/20).",
       "학습 기간과 평가 기간의 모습이 다르다: 15분 동안 가격이 전혀 안 움직인 비율 "
       "15.8% → 30.3%(20종목 모두 증가).",
       "최근 예측 모델은 과거를 한꺼번에 보는 병렬 구조(트랜스포머 · 파운데이션 모델)가 "
       "주류지만, 변동성 연구의 전통 모델(GARCH · 순환 신경망)은 과거를 차례로 읽는 순차 구조다."]
for b in BUL:
    h = th("· " + b, B, CW)
    text(LX, y, CW, h, [[("· ", dict(bold=True, color=BAR)), (b, {})]],
         size=B, color=INK, spacing=1.06)
    y += h + 0.14
y += GAP - 0.14

y += rbox(LX, y, CW, "연구 질문",
          ["어느 모델이 가장 좋은가가 아니라, 어떤 상황(예측 구간 · 변동성 국면)에서 "
           "어느 처리 방식(순차 대 병렬)이 왜 좋은가"],
          BX + 1, tsize=BX, icolor=BAR, ibold=True) + GAP

y += rbox(LX, y, CW, "연구 가설",
          ["H1. 1위 모델의 처리 방식은 예측 구간과 변동성 국면에 따라 달라진다.",
           "H2. 가장 요동칠 때 병렬 딥러닝은 1위 모델과 순환 신경망보다 오차가 크다.",
           "H3. 그 차이는 입력 형태의 차이만으로 생긴 것이 아니다.",
           [("검정: ", dict(bold=True, color=RED)),
            ("귀무가설 '두 모델(묶음)의 기대 오차가 같다'를 DM(Diebold-Mariano) 검정 + "
             "Holm 보정(유의수준 0.05)으로 기각하는지 본다.", {})]],
          BX, icolor=BAR, ibold=True, fill=LTBLUE2) + GAP

EQ = [("RV~t,H~ = √( Σ r~i~^2^ ),   r~i~ = 15분 로그수익률,  i ∈ (t, t+H]",
       "실제 변동성(예측 대상)"),
      ("QLIKE = log σ̂^2^ + RV^2^ / σ̂^2^",
       "변동성 예측 오차(작을수록 좋음, 작게 예측할수록 벌점이 큼)"),
      ("국면 Q~t~ = k  ⇔  RV~t−H,t~ ∈ (q~k−1~, q~k~],  q = 종목별 학습 기간 5분위수",
       "예측 시점에 이미 아는 직전 변동성으로 나눈 5단계"),
      ("DM = d̄ / √Var~HAC~(d̄),   d~t~ = L~A,t~ − L~B,t~",
       "두 모델 오차 차이가 우연인지 보는 검정 통계량")]
y += rbox(LX, y, CW, "필수 수식",
          [[(e, dict(color=BAR, bold=True)), ("   … " + d, dict(color=GREY, bold=False))]
           for e, d in EQ],
          BX, fill=LTBLUE) + GAP
LEFT_BOTTOM = y - GAP

# ================================================================ 오른쪽 단
y = y0 + 2.3 + 0.55
hcap = th("표 1. 데이터 요약", s_(20), CW)
text(RX, y, CW, hcap, "표 1. 데이터 요약", size=s_(20), color=BAR, bold=True)
y += hcap + CGAP
T1 = [["항목", "내용"],
      ["종목 · 기간", "업비트 20종목 · 2023-10-06 ~ 2026-10-05 (2종목은 2026-07 상장폐지)"],
      ["봉 간격 · 규모", "15분봉 · 총 2,029,138봉(종목당 약 10만)"],
      ["학습 / 평가", "시간순 분할 · 학습 ~2025-11-10 / 평가 2025-11-11~ (약 70% / 30%)"],
      ["예측 대상", "15분 · 30분 · 1시간 · 4시간 · 12시간 뒤의 실현변동성(RV)"],
      ["평가 지표", "QLIKE(낮을수록 좋음, 변동성을 작게 예측할수록 벌점이 커짐)"],
      ["반복", "학습하는 모델은 무작위 시드 5개 평균"]]
y += table(RX, y, CW, T1, [26, 74], TH_, TB, rowpad=0.38, bolds={0}) + GAP

hcap = th("표 2. 비교한 27개 모델: 과거를 읽는 방식과 입력", s_(20), CW)
text(RX, y, CW, hcap, "표 2. 비교한 27개 모델: 과거를 읽는 방식과 입력",
     size=s_(20), color=BAR, bold=True)
y += hcap + CGAP
T2 = [["처리 방식", "모델", "입력(알고리즘이 받는 형태를 따름)"],
      [("순차(차례로 읽기)", SEQ),
       "GARCH 3종(GARCH-t · MS-GARCH · TAR-GARCH) · GRU · LSTM",
       "15분봉 수익률을 한 시점씩(순환 신경망은 직전 96개 = 24시간)"],
      [("병렬 · 특징 기반", PARF),
       "LightGBM · XGBoost · HistGBM · GARCH+LightGBM · Nystroem+Ridge · KernelRidge · SVR",
       "15분봉에서 만든 여러 길이(15분~7일)의 과거 변동성 특징"],
      [("병렬 · 딥러닝(신규 15종)", PARD),
       "트랜스포머 4 · 합성곱 3 · S-Mamba · 사전학습 파운데이션 7(TimesFM · Chronos 등)",
       "예측 구간 간격으로 묶은 과거 변동성 수열(한 수열의 다음 값을 예측하는 구조라 "
       "이 형태로 넣음)"]]
y += table(RX, y, CW, T2, [23, 39, 38], TH_, TB, rowpad=0.32) + GAP

y += rbox(RX, y, CW, "용어 풀이",
          [[("· 1위 모델: ", dict(bold=True, color=BAR)),
            ("같은 예측 구간 · 같은 국면에서 평균 오차(QLIKE)가 가장 작은 모델. "
             "후보 29개(4 · 12시간 28개) 중, 학습하는 모델은 시드 5개 평균.", {})],
           [("· 통계적으로 구분 안 됨: ", dict(bold=True, color=BAR)),
            ("1위와의 오차 차이가 우연 수준(Holm 보정 후 p ≥ 0.05). "
             "성능이 같다는 증명은 아니다.", {})],
           [("· 칸: ", dict(bold=True, color=BAR)),
            ("예측 구간 5개 × 범위 6개(전체 기간 + 국면 Q1~Q5) = 30칸.", {})],
           [("· 묶음 비교: ", dict(bold=True, color=BAR)),
            ("처리 방식 묶음마다 구성원 오차를 단순 평균해 비교"
             "(잘한 모델만 골라 비교하는 편향을 피함).", {})]],
          BX, fill=LTBLUE2) + GAP
RIGHT_BOTTOM = y - GAP

# ================================================================ 결과
yr = max(LEFT_BOTTOM, RIGHT_BOTTOM) + 0.65
secbar(ML, yr, FW, "결   과")
y = yr + 2.3 + 0.50

RES = [("결과 1", "가장 요동칠 때 병렬 딥러닝은 73번 비교(모델 15종 × 예측 구간 5개, "
                  "4 · 12시간은 14종) 모두 1위 모델보다 통계적으로 나빴다.", PARD),
       ("결과 2", "순환 신경망 묶음은 병렬 딥러닝 묶음보다 30칸 중 14칸에서 나았고 "
                  "0칸에서 졌다. 같은 입력을 줘도 29칸에서 나았다.", SEQ)]

# ---- 그림 1 (왼쪽) + 오른쪽에 결과 1 · 결과 2 · 그림 1 캡션
pic("nf/f1.png", ML, y, WF1)
h1 = WF1 * ar("nf/f1.png")
cx1 = ML + WF1 + 1.3
cw1 = MR - cx1
c1 = ("예측 구간별 상위 5개 모델(위: 전체 기간, 아래: 가장 요동칠 때). 가로축 = 1위와의 "
      "오차 차이(0 = 1위), 색 = 처리 방식(파랑 순차, 주황 트리 · 커널, 갈색 병렬 딥러닝). "
      "채운 점 = 1위와 통계적으로 구분 안 됨, 빈 점 = 1위보다 유의하게 나쁨, 점선 아래 = "
      "5위 밖 병렬 딥러닝 중 가장 나은 모델. 전체 기간 5위 안은 대부분 트리 · 커널이고, "
      "가장 요동칠 때 1시간~12시간은 순환 신경망(GRU · LSTM)이 1위로 올라온다. "
      "병렬 딥러닝은 거의 모두 5위 밖이다.")
hbs = [th(l + "   " + t, B, cw1 - 1.0, spacing=1.04) + 0.42 for l, t, _ in RES]
hc1 = th("그림 1. " + c1, CAP, cw1)
slack = max(0.0, h1 - sum(hbs) - hc1)
g = min(1.3, slack / 3.0)
yy = y + (slack - 2 * g if slack > 3 * g else 0) * 0.0
for (lab, txt_, col), hb in zip(RES, hbs):
    box(cx1, yy, cw1, hb, LTBLUE2, EDGE, 1.2, MSO_SHAPE.ROUNDED_RECTANGLE, 0.10)
    text(cx1 + 0.5, yy, cw1 - 1.0, hb,
         [[(lab + "   ", dict(bold=True, color=col)), (txt_, {})]],
         size=B, color=INK, spacing=1.04, anchor=MSO_ANCHOR.MIDDLE)
    yy += hb + max(0.22, g)
text(cx1, yy, cw1, max(hc1, h1 - (yy - y)),
     [[("그림 1. ", dict(bold=True, color=BAR)), (c1, {})]],
     size=CAP, color=GREY, spacing=1.05)
y += max(h1, (yy - y) + hc1) + GAP

# ---- 그림 2 (왼쪽) + 오른쪽 옆 캡션
pic("nf/f2.png", ML, y, WF2)
h2 = WF2 * ar("nf/f2.png")
cx = ML + WF2 + 1.3; cw2 = MR - cx
c2 = ("국면(Q1 가장 잔잔 → Q5 가장 요동)별로 각 처리 방식에서 가장 나은 모델이 1위와 "
      "얼마나 차이 나는지(0이면 그 방식의 모델이 1위). 파랑 = 순차 5종 중 가장 나은 모델, "
      "주황 = 트리 · 커널 7종 중, 갈색 = 병렬 딥러닝 15종 중. 채운 점 = 그 방식에 1위와 "
      "구분 안 되는 모델이 있음. 잔잔 · 보통 국면(Q1~Q4)은 대부분 주황이 0이고(예외: "
      "가장 잔잔한 4 · 12시간은 갈색, 15분 Q3 · 12시간 Q4는 파랑), 가장 요동칠 때 "
      "1시간 · 4시간 · 12시간은 파랑이 0이다(각각 GRU · LSTM · GRU). 갈색은 요동칠수록 "
      "1위에서 멀어진다.")
hc2 = th("그림 2. " + c2, CAP, cw2)
text(cx, y, cw2, max(h2, hc2), [[("그림 2. ", dict(bold=True, color=BAR)), (c2, {})]],
     size=CAP, color=GREY, spacing=1.05, anchor=MSO_ANCHOR.MIDDLE)
y += max(h2, hc2) + GAP

# ---- 그림 3 (왼쪽) + 표 3 · 표 4 (오른쪽)
yb = y
pic("nf/f3.png", LX, yb, WF3)
h3 = WF3 * ar("nf/f3.png")
c3 = ("데이터 탐색(평가 기간 20종목). 왼쪽: 다음 변동성 ÷ 직전 변동성(로그). 잔잔할 때는 "
      "다음에 커지고 요동칠 때는 줄어든다(가장 요동칠 때 다섯 구간 모두 20/20종목 음수). "
      "오른쪽: 직전 구간 안에서 변동성이 뒤쪽 절반에 몰린 정도와 왼쪽 값의 순위 상관. "
      "가장 요동칠 때 모든 구간에서 양(+)이고, 30분 · 1시간 · 4시간은 가장 잔잔할 때보다 "
      "유의하게 크다(12시간은 구분 안 됨, 표 4). 이 '구간 안 흐름'은 15분봉을 받는 "
      "모델(순환 신경망 · 트리 · 커널 · GARCH)만 입력으로 받는다.")
hc3 = th("그림 3. " + c3, CAP, CW)
text(LX, yb + h3 + CGAP, CW, hc3,
     [[("그림 3. ", dict(bold=True, color=BAR)), (c3, {})]],
     size=CAP, color=GREY, spacing=1.05)
bot_l = yb + h3 + CGAP + hc3

yy = yb
t3c = ("처리 방식 묶음끼리의 통계 검정(30칸 중 칸 수). 귀무가설 = 두 묶음의 평균 오차가 "
       "같다, 대립가설 = 다르다. DM 검정 후 비교마다 30칸에 Holm 보정, p < 0.05면 우세. "
       "'같은 입력'은 순환 신경망에 병렬 딥러닝과 같은 변동성 수열 · 학습 표본을 준 통제 비교.")
h = th("표 3. " + t3c, CAP, CW)
text(RX, yy, CW, h, [[("표 3. ", dict(bold=True, color=BAR)), (t3c, {})]],
     size=CAP, color=GREY, spacing=1.05)
yy += h + CGAP
T3 = [["비교(A 대 B)", "A 우세", "B 우세", "구분 안 됨", "가장 요동칠 때 A 우세 구간"],
      ["순환 신경망 대 병렬 딥러닝", "14", "0", "16", "30분 · 1시간 · 4시간 · 12시간"],
      ["순차 5종 대 병렬 딥러닝", "11", "1", "18", "30분 · 1시간 · 4시간"],
      ["순차 5종 대 트리 · 커널", "1", "16", "13", "4시간"],
      ["같은 입력의 순환 신경망 대 병렬 딥러닝", "29", "0", "1", "15분~12시간 모두"]]
yy += table(RX, yy, CW, T3, [33, 11, 11, 13, 32], TH_, TB, rowpad=0.26, bolds={0},
            aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 3 + [PP_ALIGN.LEFT]) + GAP

t4c = ("그림 3 오른쪽의 검정(종목 단위). 양(+) 종목 = 가장 요동칠 때 상관이 양수인 종목 "
       "수(부호 검정), Q5 > Q1 = 같은 종목에서 요동 때 상관이 잔잔 때보다 큰 종목 수"
       "(윌콕슨 부호순위 검정, 귀무가설 = 차이의 중앙값 0, Holm 보정). 15분은 직전 구간이 "
       "15분봉 1개라 해당 없음. 종목들이 같은 시장 충격을 받아 p값은 근사.")
h = th("표 4. " + t4c, CAP, CW)
text(RX, yy, CW, h, [[("표 4. ", dict(bold=True, color=BAR)), (t4c, {})]],
     size=CAP, color=GREY, spacing=1.05)
yy += h + CGAP
T4 = [["예측 구간", "가장 요동(Q5) 상관 · 양(+) 종목", "가장 잔잔(Q1) 상관",
       "Q5 > Q1 종목 · 보정 p"],
      ["30분", "+0.075 · 19/20", "−0.007", "11/14 · 0.049"],
      ["1시간", "+0.088 · 20/20", "−0.008", "20/20 · < 0.001"],
      ["4시간", "+0.195 · 20/20", "+0.078", "17/20 · 0.001"],
      ["12시간", "+0.164 · 18/19", "+0.145", "8/16 · 0.53(구분 안 됨)"]]
yy += table(RX, yy, CW, T4, [18, 32, 22, 28], TH_, TB, rowpad=0.26, bolds={0},
            aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 3)
RESULT_BOTTOM = max(bot_l, yy)

# ================================================================ 결론
yc = RESULT_BOTTOM + 0.65
secbar(ML, yc, FW, "결   론")
y = yc + 2.3 + 0.50
CON = [("H1 지지: ", "1위 처리 방식은 상황에 따라 바뀐다. 평소엔 트리 · 커널, 가장 요동칠 때 "
                     "1~12시간엔 순환 신경망(트리 · 커널과는 구분 안 됨)."),
       ("H2 지지: ", "가장 요동칠 때 병렬 딥러닝은 73번 비교 모두 1위보다 나빴고, 순환 신경망 "
                     "묶음보다도 30분~12시간에서 나빴다(15분은 구분 안 됨)."),
       ("H3 부분 지지: ", "같은 입력을 줘도 순환 신경망 묶음이 30칸 중 29칸 우세. 다만 요동 때 "
                          "1위 자리는 15분봉 입력 순환 신경망에만 해당한다.")]
for i, (lab, bd) in enumerate(CON):
    hh = th(lab + bd, B, FW - 2.0, spacing=1.02)
    box(ML, y + 0.02, 1.0, 1.0, BAR, shape=MSO_SHAPE.OVAL)
    text(ML, y + 0.02, 1.0, 1.0, str(i + 1), size=s_(16), color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(ML + 1.5, y, FW - 1.5, hh,
         [[(lab, dict(bold=True, color=RED)), (bd, {})]],
         size=B, color=INK, spacing=1.02)
    y += max(hh, 1.05) + 0.12
y += GAP - 0.12

t5c = ("왜 그 모델이 1위인가: 결과 → 데이터 탐색 근거 → 문헌. 결과와 데이터 탐색의 방향이 "
       "같다는 해석이며 인과 증명은 아니다. 문헌은 주가지수 일 단위 자료라 암호화폐 "
       "15분~12시간으로 옮기는 것은 유추다.")
h = th("표 5. " + t5c, CAP, FW)
text(ML, y, FW, h, [[("표 5. ", dict(bold=True, color=BAR)), (t5c, {})]],
     size=CAP, color=GREY, spacing=1.05)
y += h + CGAP
T5 = [["상황", "1위", "데이터 탐색 근거(그림 3 · 표 4)", "뒷받침 문헌"],
      ["잔잔 · 보통(Q1~Q4)", ("대부분 트리 · 커널", PARF),
       "잔잔할 때는 구간 안 흐름 신호가 약하다(30분 · 1시간 Q1 상관 −0.01). 직전 변동성의 "
       "수준이 주된 정보이고, 트리 · 커널은 여러 길이의 변동성 수준을 특징으로 받는다",
       "Corsi(2009): 변동성은 짧은 · 중간 · 긴 시간 척도 성분으로 잘 예측된다"],
      ["가장 요동(Q5), 1시간~12시간", ("순환 신경망(트리 · 커널과 구분 안 됨)", SEQ),
       "직전 구간 안에서 변동성이 뒤쪽에 몰렸는지가 다음 변동성과 연결된다(Q5 상관 "
       "+0.09~+0.20). 15분봉을 받는 모델만 이 흐름을 입력으로 받는다",
       "Bucci(2020): 순환 신경망(LSTM 등)이 실현변동성 예측에서 계량 모형과 경쟁하거나 더 낫다"],
      ["가장 요동(Q5), 15분 · 30분", ("커널(Nystroem+Ridge)", PARF),
       "15분은 직전 구간이 봉 1개라 구간 안 흐름이 없고, 30분은 신호가 약하다(+0.075)",
       "해당 없음(이 짧은 구간을 다룬 문헌을 확인하지 못함)"],
      ["병렬 딥러닝", ("가장 잔잔한 4 · 12시간만", PARD),
       "이번 입력(변동성 수열)에는 구간 안 흐름이 없다. 단 같은 입력의 순환 신경망도 "
       "29/30칸 앞서 입력만으로는 설명되지 않는다",
       "Zeng 외(2023): 시계열 트랜스포머가 단순한 선형 모형을 꾸준히 이기지 못한다"],
      ["반례: GARCH 3종", ("요동 4 · 12시간 열세", SEQ),
       "같은 15분봉을 받고도 졌다. 구간 안 흐름은 여러 요인 중 하나다",
       "Lamoureux · Lastrapes(1990): 구조 변화가 있으면 GARCH의 지속성이 과대추정된다"]]
y += table(ML, y, FW, T5, [16, 19, 38, 27], TH_, TB, rowpad=0.24, bolds={0}) + GAP

LIM = ["학습 방식(학습률 탐색 · 조기 종료 · 재학습)을 모델마다 맞추지 않아 순차 구조 자체의 "
       "효과로 단정할 수 없다.",
       "입력은 알고리즘이 받는 형태를 따랐다. 같은 변동성 수열을 받은 순환 신경망은 요동 때 "
       "1 · 4시간에서 1위보다 유의하게 나빴다.",
       "학습하는 모델은 시드 5개라 요동 때 1위가 시드마다 바뀐다(4시간 LSTM 3/5, "
       "12시간 GRU 1/5).",
       "12시간은 국면당 평가 관측이 약 2,000~3,000개로 적고, 15분 가장 잔잔한 국면은 "
       "5종목에서 가격 정지와 섞인다."]
NXT = ["학습 방식을 맞춘 통제 비교로 구조의 효과를 분리한다.",
       "시드를 10개 이상으로 늘려 요동 때 4 · 12시간 1위의 안정성을 확인한다.",
       "종목을 묶은 패널 검정으로 12시간의 표본 부족을 보완한다.",
       "직전 변동성 국면(예측 시점에 이미 아는 값)에 따라 모델을 골라 쓰는 방식을 별도 "
       "평가 기간에서 검증한다."]
def mk(items):
    return [[("· ", dict(bold=True, color=BAR)), (t, {})] for t in items]
hL = rbox(LX, y, CW, "한계", mk(LIM), B, tsize=B, pad=0.40, gap=0.08,
          fill=LTBLUE2, spacing=1.02)
hR = rbox(RX, y, CW, "다음 연구", mk(NXT), B, tsize=B, pad=0.40, gap=0.08,
          fill=LTBLUE2, spacing=1.02)
y += max(hL, hR) + 0.30
ack = ("이 성과는 정부(과학기술정보통신부)의 재원으로 한국연구재단의 지원을 받아 수행된 "
       "연구임 (RS-2025-00558517).")
hak = th(ack, CAP, FW)
text(ML, y, FW, hak, ack, size=CAP, color=GREY, align=PP_ALIGN.CENTER)
FINAL = y + hak

print(f"좌={LEFT_BOTTOM:.1f} 우={RIGHT_BOTTOM:.1f} | 결과바={yr:.1f} 결과끝={RESULT_BOTTOM:.1f}"
      f" | 결론바={yc:.1f} 최종={FINAL:.1f} 여유={PANEL_B-FINAL:.2f}")
if os.environ.get("SAVE") == "1": prs.save("poster3.pptx")
